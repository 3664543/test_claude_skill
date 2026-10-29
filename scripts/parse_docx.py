# -*- coding: utf-8 -*-
"""Parse a .docx requirement document into structured JSON + image files.

python-docx covers paragraphs and tables; textboxes need an xpath into the
raw XML; images are pulled off document relationships and saved to out_dir.
stdout is ASCII only (Chinese never prints cleanly on a Windows console).
"""
import json
import os
import sys

from docx import Document

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"


def extract(docx_path, out_dir):
    os.makedirs(out_dir, exist_ok=True)
    doc = Document(docx_path)

    paragraphs = [{"style": p.style.name if p.style is not None else "",
                   "text": p.text.strip()}
                  for p in doc.paragraphs if p.text.strip()]

    tables = []
    for tbl in doc.tables:
        tables.append([[c.text.strip() for c in row.cells] for row in tbl.rows])

    # textboxes (VML + DrawingML): every w:txbxContent's inner text. Word 2010+
    # writes each textbox twice (mc:Choice DrawingML + mc:Fallback VML) — keep
    # only the Choice branch, bare VML boxes have no Fallback ancestor.
    textboxes = []
    for txbx in doc.element.body.iter(W_NS + "txbxContent"):
        if any(a.tag == MC_FALLBACK for a in txbx.iterancestors()):
            continue
        text = "".join(t.text or "" for t in txbx.iter(W_NS + "t")).strip()
        if text:
            textboxes.append(text)

    images = []
    for rel in doc.part.rels.values():
        if "image" not in rel.reltype:
            continue
        part = rel.target_part
        ext = os.path.splitext(str(part.partname))[1] or ".png"
        idx = len(images) + 1
        img_path = os.path.abspath(os.path.join(out_dir, "image%d%s" % (idx, ext)))
        with open(img_path, "wb") as f:
            f.write(part.blob)
        images.append(img_path)

    result = {"docx": os.path.abspath(docx_path), "out_dir": os.path.abspath(out_dir),
              "paragraphs": paragraphs, "tables": tables,
              "textboxes": textboxes, "images": images}

    with open(os.path.join(out_dir, "content.json"), "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=1)
    return result


def main(argv):
    if len(argv) != 3:
        print("usage: parse_docx.py <docx> <out_dir>")
        return 2
    result = extract(argv[1], argv[2])
    total = len(result["paragraphs"]) + len(result["tables"]) + \
        len(result["textboxes"]) + len(result["images"])
    if total == 0:
        print("EMPTY DOCUMENT: no paragraphs, tables, textboxes or images")
        return 2
    print("PARSED paragraphs=%d tables=%d textboxes=%d images=%d out=%s"
          % (len(result["paragraphs"]), len(result["tables"]),
             len(result["textboxes"]), len(result["images"]),
             result["out_dir"].encode("ascii", "replace").decode("ascii")))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
