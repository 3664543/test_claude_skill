# -*- coding: utf-8 -*-
"""Build a sample requirement docx for tests and the E2E smoke run."""
import base64
import io
import os

from docx import Document
from docx.oxml import parse_xml

# 1x1 transparent PNG
PNG_B64 = ("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUl"
           "EQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg==")

TEXTBOX_XML = (
    '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:v="urn:schemas-microsoft-com:vml">'
    '<w:r><w:pict>'
    '<v:shape style="width:120pt;height:40pt">'
    '<v:textbox><w:txbxContent>'
    '<w:p><w:r><w:t>验收标准：登录耗时小于2秒</w:t></w:r></w:p>'
    '</w:txbxContent></v:textbox>'
    '</v:shape>'
    '</w:pict></w:r></w:p>'
)

# Modern Word textboxes ship as mc:AlternateContent with a DrawingML Choice
# AND a VML Fallback holding the SAME w:txbxContent — parser must keep one.
MODERN_TEXTBOX_XML = (
    '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:v="urn:schemas-microsoft-com:vml" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
    '<w:r><mc:AlternateContent>'
    '<mc:Choice Requires="wps"><w:drawing>'
    '<wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" '
    'relativeHeight="0" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
    '<wp:simplePos x="0" y="0"/>'
    '<wp:positionH relativeFrom="column"><wp:posOffset>0</wp:posOffset></wp:positionH>'
    '<wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV>'
    '<wp:extent cx="1524000" cy="508000"/><wp:wrapNone/>'
    '<wp:docPr id="1" name="TextBox 1"/>'
    '<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
    '<wps:wsp><wps:txbx><w:txbxContent>'
    '<w:p><w:r><w:t>验收标准二：验证码发送耗时小于10秒</w:t></w:r></w:p>'
    '</w:txbxContent></wps:txbx><wps:bodyPr/></wps:wsp>'
    '</a:graphicData></a:graphic>'
    '</wp:anchor></w:drawing></mc:Choice>'
    '<mc:Fallback><w:pict>'
    '<v:shape style="width:120pt;height:40pt">'
    '<v:textbox><w:txbxContent>'
    '<w:p><w:r><w:t>验收标准二：验证码发送耗时小于10秒</w:t></w:r></w:p>'
    '</w:txbxContent></v:textbox>'
    '</v:shape></w:pict></mc:Fallback>'
    '</mc:AlternateContent></w:r></w:p>'
)


def build_fixture(out_path):
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    doc = Document()
    doc.add_heading("示例登录需求", level=1)
    doc.add_paragraph("用户使用手机号和密码登录系统。")
    doc.add_paragraph("密码连续错误5次后账号锁定30分钟。")
    doc.add_paragraph("支持短信验证码登录方式。")
    doc.add_picture(io.BytesIO(base64.b64decode(PNG_B64)))
    tbl = doc.add_table(rows=3, cols=3)
    data = [["登录方式", "说明", "备注"],
            ["密码登录", "手机号+密码", "默认"],
            ["验证码登录", "手机号+短信验证码", "60秒有效期"]]
    for i, row in enumerate(data):
        for j, text in enumerate(row):
            tbl.cell(i, j).text = text
    doc.element.body.append(parse_xml(TEXTBOX_XML))
    doc.element.body.append(parse_xml(MODERN_TEXTBOX_XML))
    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    print("fixture at:", build_fixture(os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "_out", "sample_req.docx")))
